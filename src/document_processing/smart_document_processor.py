import csv
import io
import os
import re
import tempfile
import time
from dataclasses import dataclass, field
from typing import Any, Callable, Dict, List, Optional, TYPE_CHECKING

import boto3
from botocore.exceptions import ClientError
from langchain_core.documents import Document

if TYPE_CHECKING:
    from src.langchain_components.vector_store import VectorStore


@dataclass
class ProcessingMetrics:
    total_files: int = 0
    successful: int = 0
    failed: int = 0
    errors: List[str] = field(default_factory=list)
    processing_time: float = 0.0


class SmartDocumentProcessor:
    def __init__(self, vector_store: Optional["VectorStore"] = None) -> None:
        self._vector_store = vector_store
        self._s3 = boto3.client("s3")
        self._handlers: Dict[str, Callable[[bytes, str], List[Document]]] = {
            ".pdf": self._handle_pdf,
            ".csv": self._handle_csv,
            ".xlsx": self._handle_excel,
            ".xls": self._handle_excel,
            ".docx": self._handle_word,
            ".doc": self._handle_word,
            ".txt": self._handle_text,
            ".md": self._handle_text,
            ".json": self._handle_text,
            ".png": self._handle_image,
            ".jpg": self._handle_image,
            ".jpeg": self._handle_image,
            ".tiff": self._handle_image,
            ".bmp": self._handle_image,
        }

    @staticmethod
    def _handle_pdf(content: bytes, source: str) -> List[Document]:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        docs = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text.strip():
                docs.append(Document(
                    page_content=text,
                    metadata={"source": source, "page": i + 1, "type": "pdf"},
                ))
        return docs

    @staticmethod
    def _handle_csv(content: bytes, source: str) -> List[Document]:
        text = content.decode("utf-8", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        docs = []
        for i, row in enumerate(reader):
            row_text = "\n".join(f"{k}: {v}" for k, v in row.items() if v)
            if row_text.strip():
                docs.append(Document(
                    page_content=row_text,
                    metadata={"source": source, "row": i + 1, "type": "csv"},
                ))
        return docs

    @staticmethod
    def _handle_excel(content: bytes, source: str) -> List[Document]:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        docs = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_text = []
            for row in ws.iter_rows(values_only=True):
                line = " | ".join(str(c) for c in row if c is not None)
                if line.strip():
                    rows_text.append(line)
            if rows_text:
                docs.append(Document(
                    page_content="\n".join(rows_text),
                    metadata={"source": source, "sheet": sheet_name, "type": "excel"},
                ))
        wb.close()
        return docs

    @staticmethod
    def _handle_word(content: bytes, source: str) -> List[Document]:
        from docx import Document as DocxDocument

        doc = DocxDocument(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        if paragraphs:
            return [Document(
                page_content="\n".join(paragraphs),
                metadata={"source": source, "type": "word"},
            )]
        return []

    @staticmethod
    def _chunk_text(text: str, chunk_size: int = 4000, overlap: int = 200) -> List[str]:
        if len(text) <= chunk_size:
            return [text]
        chunks: List[str] = []
        start = 0
        while start < len(text):
            end = min(start + chunk_size, len(text))
            if end < len(text):
                last_space = text.rfind(" ", start + chunk_size // 2, end)
                if last_space > start:
                    end = last_space
            chunks.append(text[start:end].strip())
            if end >= len(text):
                break
            start = end - overlap
            if start < 0:
                start = 0
        return [c for c in chunks if c]

    @staticmethod
    def _handle_text(content: bytes, source: str) -> List[Document]:
        text = content.decode("utf-8", errors="replace")
        chunks = SmartDocumentProcessor._chunk_text(text.strip())
        doc_type = "markdown" if source.lower().endswith(".md") else "text"
        return [
            Document(
                page_content=chunk,
                metadata={"source": source, "type": doc_type, "chunk": i},
            )
            for i, chunk in enumerate(chunks)
        ]

    @staticmethod
    def _handle_image(content: bytes, source: str) -> List[Document]:
        try:
            from PIL import Image
            import pytesseract

            image = Image.open(io.BytesIO(content))
            text = pytesseract.image_to_string(image)
            if text.strip():
                return [Document(
                    page_content=text,
                    metadata={"source": source, "type": "image"},
                )]
        except Exception:
            pass
        return []

    def _detect_handler(self, filename: str) -> Optional[Callable[[bytes, str], List[Document]]]:
        _, ext = os.path.splitext(filename)
        return self._handlers.get(ext.lower())

    def process_file(self, content: bytes, filename: str, source: str = "") -> List[Document]:
        handler = self._detect_handler(filename)
        if handler is None:
            raise ValueError(f"No handler for file type: {filename}")
        return handler(content, source or filename)

    def process_from_s3(self, bucket: str, key: str) -> List[Document]:
        response = self._s3.get_object(Bucket=bucket, Key=key)
        content = response["Body"].read()
        return self.process_file(content, os.path.basename(key), source=f"s3://{bucket}/{key}")

    def process_batch_from_s3(self, bucket: str, prefix: str = "") -> ProcessingMetrics:
        metrics = ProcessingMetrics()
        start = time.time()

        if self._vector_store is None:
            from src.langchain_components.vector_store import VectorStore
            self._vector_store = VectorStore()

        paginator = self._s3.get_paginator("list_objects_v2")
        pages = paginator.paginate(Bucket=bucket, Prefix=prefix)

        for page in pages:
            for obj in page.get("Contents", []):
                key = obj["Key"]
                metrics.total_files += 1
                try:
                    docs = self.process_from_s3(bucket, key)
                    if docs:
                        self._vector_store.add_documents([
                            {"text": d.page_content, "metadata": d.metadata}
                            for d in docs
                        ])
                    metrics.successful += 1
                except Exception as exc:
                    metrics.failed += 1
                    metrics.errors.append(f"{key}: {exc}")

        metrics.processing_time = time.time() - start
        return metrics

    def process_local_file(self, filepath: str) -> List[Document]:
        with open(filepath, "rb") as f:
            content = f.read()
        return self.process_file(content, os.path.basename(filepath), source=filepath)
